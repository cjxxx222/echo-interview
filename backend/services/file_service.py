import os
import uuid
from fastapi import UploadFile
from config import UPLOAD_DIR


def save_upload_file(file: UploadFile) -> tuple[str, str]:
    """保存上传文件，返回 (resume_id, file_path)"""
    os.makedirs(UPLOAD_DIR, exist_ok=True)
    resume_id = str(uuid.uuid4())
    ext = os.path.splitext(file.filename or "")[1].lower()
    file_path = os.path.join(UPLOAD_DIR, f"{resume_id}{ext}")
    with open(file_path, "wb") as f:
        content = file.file.read()
        f.write(content)
    return resume_id, file_path


def extract_text(file_path: str) -> str:
    """根据文件后缀提取文本内容"""
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(file_path)
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    elif ext == ".docx":
        from docx import Document
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        doc = Document(file_path)
        # 简历模板大量使用表格排版，而 doc.paragraphs 拿不到表格内文字；
        # 若把 doc.paragraphs 与 doc.tables 分开取，又会打乱原有顺序
        # （表格内容被统一追加到末尾）。这里按 body 的真实顺序遍历两者。
        lines: list[str] = []
        for child in doc.element.body.iterchildren():
            if child.tag == qn("w:p"):
                lines.append(Paragraph(child, doc).text)
            elif child.tag == qn("w:tbl"):
                for row in Table(child, doc).rows:
                    cells = [c.text.strip() for c in row.cells]
                    line = " | ".join(c for c in cells if c)
                    if line:
                        lines.append(line)
        return "\n".join(lines)

    elif ext in (".txt", ".md"):
        # 中文用户的 txt 可能是 GBK/GB18030（Windows 记事本旧默认编码），
        # 直接按 utf-8 解码会抛 UnicodeDecodeError → 上传直接 500。
        raw = open(file_path, "rb").read()
        for enc in ("utf-8", "gb18030"):
            try:
                return raw.decode(enc)
            except UnicodeDecodeError:
                continue
        return raw.decode("utf-8", errors="ignore")

    else:
        raise ValueError(f"不支持的文件格式：{ext}")
