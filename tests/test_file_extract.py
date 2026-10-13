"""简历文本提取测试：覆盖真实简历最常见的两种坑（表格排版 / GBK 编码）。"""
import pytest

from services.file_service import extract_text


def _make_docx(path, with_table: bool = True):
    import docx

    d = docx.Document()
    d.add_paragraph("林晓雨")
    d.add_paragraph("求职意向：AI 产品经理（校招）")
    if with_table:
        t = d.add_table(rows=2, cols=2)
        t.rows[0].cells[0].text = "技能"
        t.rows[0].cells[1].text = "SQL、Axure、Prompt"
        t.rows[1].cells[0].text = "实习"
        t.rows[1].cells[1].text = "某互联网公司 AI 产品部"
    d.save(path)


def test_docx_extracts_paragraphs(tmp_path):
    p = tmp_path / "resume.docx"
    _make_docx(p)
    text = extract_text(str(p))
    assert "林晓雨" in text
    assert "AI 产品经理" in text


def test_docx_extracts_table_text(tmp_path):
    """回归：简历常用表格排版，doc.paragraphs 取不到表格内文字，内容会大面积丢失。"""
    p = tmp_path / "resume.docx"
    _make_docx(p)
    text = extract_text(str(p))
    assert "SQL" in text and "Axure" in text, "表格内的技能信息不能丢"
    assert "某互联网公司 AI 产品部" in text, "表格内的实习信息不能丢"


def test_docx_preserves_document_order(tmp_path):
    """回归：段落与表格必须按文档真实顺序输出，而不是把表格全堆到末尾。"""
    import docx

    p = tmp_path / "order.docx"
    d = docx.Document()
    d.add_paragraph("第一部分-教育")
    t = d.add_table(rows=1, cols=1)
    t.rows[0].cells[0].text = "表格-复旦大学"
    d.add_paragraph("第二部分-技能")
    d.save(p)

    text = extract_text(str(p))
    assert text.index("第一部分-教育") < text.index("表格-复旦大学") < text.index("第二部分-技能"), \
        "表格内容应保持在其原始位置"


def test_txt_utf8(tmp_path):
    p = tmp_path / "resume.txt"
    p.write_bytes("姓名：林晓雨\n技能：SQL".encode("utf-8"))
    assert "林晓雨" in extract_text(str(p))


def test_txt_gbk_does_not_crash(tmp_path):
    """回归：中文用户的 txt 可能是 GBK/GB18030，按 utf-8 硬解会 500。"""
    p = tmp_path / "resume.txt"
    p.write_bytes("姓名：林晓雨\n技能：SQL、Axure".encode("gb18030"))
    text = extract_text(str(p))
    assert "林晓雨" in text and "Axure" in text


def test_md_uses_same_path_as_txt(tmp_path):
    p = tmp_path / "resume.md"
    p.write_bytes("# 林晓雨\n\n技能：SQL".encode("utf-8"))
    assert "林晓雨" in extract_text(str(p))


def test_unsupported_ext_raises(tmp_path):
    p = tmp_path / "resume.rtf"
    p.write_text("x", encoding="utf-8")
    with pytest.raises(ValueError):
        extract_text(str(p))
